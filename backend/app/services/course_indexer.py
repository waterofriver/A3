import hashlib
import json
import re
from collections import defaultdict
from pathlib import Path
from urllib.parse import quote
from xml.etree import ElementTree
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader

from app.db.database import Database
from app.repositories.courses import CourseRepository
from app.schemas.course import (
    CourseBaseData,
    CourseChapterData,
    CourseDocumentData,
)

MAX_PREVIEW_CHARS = 200_000
SLUG_PATTERN = re.compile(r"^[a-z0-9][a-z0-9-]{0,63}$")
TEXT_EXTENSIONS = {
    ".c",
    ".cpp",
    ".csv",
    ".go",
    ".h",
    ".hpp",
    ".ini",
    ".java",
    ".js",
    ".json",
    ".jsx",
    ".md",
    ".ps1",
    ".py",
    ".rs",
    ".sh",
    ".toml",
    ".ts",
    ".tsx",
    ".txt",
    ".yaml",
    ".yml",
}
IMAGE_EXTENSIONS = {".bmp", ".gif", ".jpeg", ".jpg", ".png", ".webp"}
VIDEO_EXTENSIONS = {".m4v", ".mov", ".mp4", ".webm"}
SUPPORTED_EXTENSIONS = (
    TEXT_EXTENSIONS
    | IMAGE_EXTENSIONS
    | VIDEO_EXTENSIONS
    | {".pdf", ".docx", ".pptx"}
)


def resolve_under(root: Path, relative_path: Path) -> Path:
    resolved_root = root.resolve()
    resolved = (resolved_root / relative_path).resolve()
    if not resolved.is_relative_to(resolved_root):
        raise ValueError("文档路径必须位于课程根目录内。")
    return resolved


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def extract_preview(path: Path, extension: str) -> str | None:
    if extension in TEXT_EXTENSIONS:
        return path.read_text(encoding="utf-8", errors="replace")[:MAX_PREVIEW_CHARS]
    if extension == ".pdf":
        reader = PdfReader(path)
        return "\n\n".join(
            (page.extract_text() or "") for page in reader.pages[:30]
        )[:MAX_PREVIEW_CHARS]
    if extension == ".docx":
        document = Document(path)
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            parts.extend(
                "\t".join(cell.text for cell in row.cells) for row in table.rows
            )
        return "\n".join(part for part in parts if part.strip())[:MAX_PREVIEW_CHARS]
    if extension == ".pptx":
        with ZipFile(path) as archive:
            slide_names = sorted(
                (
                    name
                    for name in archive.namelist()
                    if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
                ),
                key=lambda name: int(re.search(r"\d+", name).group()),
            )
            parts = []
            for slide_name in slide_names:
                root = ElementTree.fromstring(archive.read(slide_name))
                parts.extend(
                    node.text.strip()
                    for node in root.iter()
                    if node.tag.endswith("}t") and node.text and node.text.strip()
                )
        return "\n\n".join(parts)[:MAX_PREVIEW_CHARS]
    return None


def course_data(course, documents) -> CourseBaseData:
    chapters: dict[tuple[str, str], list[CourseDocumentData]] = defaultdict(list)
    for document in documents:
        chapters[(document.chapter_path, document.chapter_name)].append(
            CourseDocumentData(
                id=document.id,
                chapter_path=document.chapter_path,
                filename=document.filename,
                file_type=document.file_type,
                sha256=document.sha256,
                size_bytes=document.size_bytes,
                preview_text=document.preview_text,
                media_url=document.media_url,
            )
        )
    return CourseBaseData(
        name=course.name,
        slug=course.slug,
        content_ready=course.content_ready,
        is_demo=course.is_demo,
        chapters=[
            CourseChapterData(path=path, name=name, documents=items)
            for (path, name), items in chapters.items()
        ],
    )


class CourseIndexer:
    def __init__(self, db: Database):
        self.db = db

    def index_root(self, root: Path) -> list[CourseBaseData]:
        resolved_root = root.resolve()
        resolved_root.mkdir(parents=True, exist_ok=True)
        indexed_slugs: set[str] = set()

        with self.db.session() as session:
            repository = CourseRepository(session)
            for course_dir in sorted(
                (entry for entry in resolved_root.iterdir() if entry.is_dir()),
                key=lambda entry: entry.name,
            ):
                metadata_path = course_dir / "course.json"
                if not metadata_path.is_file():
                    continue
                try:
                    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
                except (OSError, UnicodeError, json.JSONDecodeError):
                    continue
                name = str(metadata.get("name", "")).strip()
                slug = str(metadata.get("slug", "")).strip()
                if not name or not SLUG_PATTERN.fullmatch(slug):
                    continue

                document_drafts = []
                for path in sorted(course_dir.rglob("*")):
                    if not path.is_file() or path.name == "course.json":
                        continue
                    extension = path.suffix.lower()
                    if extension not in SUPPORTED_EXTENSIONS:
                        continue
                    relative = path.relative_to(course_dir)
                    try:
                        safe_path = resolve_under(course_dir, relative)
                    except ValueError:
                        continue
                    chapter_path = relative.parent.as_posix()
                    chapter_name = (
                        relative.parent.name
                        if chapter_path != "."
                        else "课程资料"
                    )
                    try:
                        preview_text = extract_preview(safe_path, extension)
                    # A broken source document remains downloadable even if its preview cannot be extracted.
                    except Exception:
                        preview_text = None
                    relative_url = quote(relative.as_posix(), safe="/")
                    document_drafts.append(
                        {
                            "id": hashlib.sha256(
                                f"{slug}:{relative.as_posix()}".encode("utf-8")
                            ).hexdigest(),
                            "chapter_path": chapter_path,
                            "chapter_name": chapter_name,
                            "relative_path": relative.as_posix(),
                            "filename": path.name,
                            "file_type": extension.removeprefix("."),
                            "sha256": file_sha256(safe_path),
                            "size_bytes": safe_path.stat().st_size,
                            "preview_text": preview_text,
                            "media_url": f"/media/courses/{quote(slug)}/{relative_url}",
                        }
                    )

                course = repository.upsert_course(
                    slug=slug,
                    name=name,
                    root_path=str(course_dir.resolve()),
                    content_ready=bool(document_drafts),
                )
                repository.replace_documents(slug, document_drafts)
                indexed_slugs.add(slug)

            repository.remove_courses_not_in(indexed_slugs)
            return [
                course_data(course, repository.list_documents(course.slug))
                for course in repository.list_courses()
            ]
